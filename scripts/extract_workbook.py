#!/usr/bin/env python3
"""Extract the ASD "Essentials for Enterprise IT - Assessment Workbook" into JSON.

Usage:
    pip install openpyxl
    python3 scripts/extract_workbook.py "path/to/Assessment Workbook.xlsx" server/data/framework.json

Re-run this when ASD publishes a new version of the workbook. The app only reads
the generated JSON, so saved assessments keep working as long as action and
indicator IDs stay the same.
"""
import json
import re
import sys

import openpyxl

CAPABILITY = ["Governance and culture", "Visibility and response"]
TECHNICAL = [
    "Boundary services",
    "Identity services",
    "Privileged environments",
    "Critical processes and services",
    "Business systems",
    "User devices",
]
STRATEGIES = [
    "Patch applications",
    "Patch operating systems",
    "Multi-factor authentication",
    "Restrict administrative privileges",
    "Application control",
    "Restrict Microsoft Office macros",
    "User application hardening",
    "Regular backups",
]

PRINCIPLE_RE = re.compile(r"^Principle\s+([A-Z]{3}-\d+)\s*-\s*(.+)$")
ACTION_RE = re.compile(r"^IT\.[A-Z]{3}-\d+\.\d+$")
IGP_RE = re.compile(r"^IT\.[A-Z]{3}-\d+\.\d+\.(NE|PE|FE)-\d+$")
ISM_RE = re.compile(r"^ISM-\d+$")
LEVEL_MAP = {"Not effective": "NE", "Partially effective": "PE", "Fully effective": "FE"}


def text(v):
    if v is None:
        return ""
    return str(v).strip()


def parse_domain(ws, name, priority, description):
    domain = {"name": name, "priority": priority, "description": description, "principles": []}
    principle = None
    action = None
    in_e8 = False
    for row in range(1, ws.max_row + 1):
        b = text(ws.cell(row, 2).value)
        c = text(ws.cell(row, 3).value)
        d = text(ws.cell(row, 4).value)
        e = text(ws.cell(row, 5).value)

        m = PRINCIPLE_RE.match(b)
        if m:
            principle = {
                "id": m.group(1),
                "name": m.group(2).strip(),
                "description": text(ws.cell(row + 1, 2).value),
                "actions": [],
            }
            domain["principles"].append(principle)
            action = None
            in_e8 = False
            continue

        if ACTION_RE.match(b):
            action = {"id": b, "description": c, "indicators": [], "e8Controls": []}
            principle["actions"].append(action)
            in_e8 = False
            continue

        if action is None:
            continue

        if c == "Relevant E8 controls":
            in_e8 = True
            continue

        if IGP_RE.match(d):
            level = IGP_RE.match(d).group(1)
            action["indicators"].append({"id": d, "level": level, "text": e})
            continue

        if in_e8 and ISM_RE.match(c):
            action["e8Controls"].append(c)
            continue

        if b.startswith("Principle rating"):
            action = None
            in_e8 = False

    return domain


def parse_e8(ws):
    # Locate the E8Controls table header row.
    header_row = None
    for row in range(1, 40):
        if text(ws.cell(row, 12).value) == "ISM Control":
            header_row = row
            break
    if header_row is None:
        raise SystemExit("Could not find E8 controls table")
    headers = [text(ws.cell(header_row, col).value) for col in range(1, ws.max_column + 1)]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}

    def flag(row, col_name):
        return text(ws.cell(row, idx[col_name]).value).lower() == "true"

    controls = []
    for row in range(header_row + 1, ws.max_row + 1):
        ism = text(ws.cell(row, idx["ISM Control"]).value)
        if not ISM_RE.match(ism):
            continue
        controls.append(
            {
                "id": ism,
                "description": text(ws.cell(row, idx["Description"]).value),
                "levels": [lvl for lvl, col in (("ML1", "Maturity Level 1"), ("ML2", "Maturity Level 2"), ("ML3", "Maturity Level 3")) if flag(row, col)],
                "strategies": [s for s in STRATEGIES if flag(row, s)],
            }
        )
    return controls


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    src, dest = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src)

    # Domain priority and descriptions come from the "How to use" sheet.
    meta = {}
    how = wb["How to use"]
    for row in range(1, how.max_row + 1):
        name = text(how.cell(row, 3).value)
        if name in CAPABILITY + TECHNICAL:
            meta[name] = (text(how.cell(row, 4).value), text(how.cell(row, 5).value))

    domains = []
    for name in CAPABILITY + TECHNICAL:
        priority, description = meta.get(name, ("", ""))
        d = parse_domain(wb[name], name, priority, description)
        d["kind"] = "capability" if name in CAPABILITY else "technical"
        domains.append(d)

    # Principle ordering for the heatmap (Principle Summary sheet).
    ps = wb["Principle Summary"]
    cap_order, tech_order = [], []
    for row in range(7, ps.max_row + 1):
        a, e = text(ps.cell(row, 1).value), text(ps.cell(row, 5).value)
        if re.match(r"^[A-Z]{3}-\d+$", a):
            cap_order.append(a)
        if re.match(r"^[A-Z]{3}-\d+$", e):
            tech_order.append(e)

    controls = parse_e8(wb["Existing E8 Controls"])
    known = {c["id"] for c in controls}

    # Sanity checks
    actions = sum(len(p["actions"]) for d in domains for p in d["principles"])
    igps = sum(len(a["indicators"]) for d in domains for p in d["principles"] for a in p["actions"])
    missing = {i for d in domains for p in d["principles"] for a in p["actions"] for i in a["e8Controls"]} - known
    if missing:
        print("WARNING: mapped ISM controls missing from E8 table:", sorted(missing))

    out = {
        "source": src.split("/")[-1],
        "domains": domains,
        "principleOrder": {"capability": cap_order, "technical": tech_order},
        "e8": {"strategies": STRATEGIES, "controls": controls},
    }
    with open(dest, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    print(f"Domains: {len(domains)}, actions: {actions}, indicators: {igps}, E8 controls: {len(controls)}")


if __name__ == "__main__":
    main()
